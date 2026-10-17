"""Fabrique un jeu de données ENTIÈREMENT FICTIF pour essayer l'application.

Foyer « Famille Exemple » ; comptes, montants, commerçants et numéros inventés.
Formats identiques aux exports réels (LCL, Amex, factures et relevés PDF, classeur d'inventaire),
pour éprouver les importeurs sans jamais toucher de données réelles.

Usage : python outils/donnees_fictives.py   → dossier donnees-fictives/
"""
import pathlib, random, datetime as dt, math
from openpyxl import Workbook
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas
from PIL import Image, ImageDraw, ImageFont

RACINE = pathlib.Path(__file__).resolve().parent.parent
SORTIE = RACINE / "donnees-fictives"
SORTIE.mkdir(exist_ok=True)
random.seed(17)
D = lambda s: dt.date.fromisoformat(s)
fr = lambda d: d.strftime("%d/%m/%Y")
def mt(v):  # montant à la française, sans séparateur de milliers (comme l'export LCL)
    s = f"{v:.2f}".replace(".", ",")
    return s[:-3] if s.endswith(",00") else s

COMPTE = "09999 000099Z"
# ---------------------------------------------------------------- achats par carte
COMMERCANTS = [  # libellé bancaire, MCC, catégorie LCL, (min, max)
    ("BOUCHERIE DU MARCHE VILLE", "5422", "Alimentation:Epicerie", (18, 60)),
    ("PRIMEUR DE LA PLACE VILLE", "5499", "Alimentation:Epicerie", (6, 30)),
    ("SUPERMARCHE EXEMPLE CAMPAGNE", "5411", "Alimentation", (40, 160)),
    ("PHARMACIE CENTRALE VILLE", "5912", "Sante:Pharmacie", (8, 45)),
    ("LIBRAIRIE DES ECOLES VILLE", "5942", "Loisirs/culture", (12, 50)),
    ("BRASSERIE DU COIN VILLE", "5812", "Alimentation:Restaurant", (35, 120)),
    ("STATION SERVICE CAMPAGNE", "5541", "Automobile", (40, 90)),
    ("FROMAGERIE PERRET VILLE", "5499", "Alimentation:Epicerie", (10, 35)),
]
def achats(debut, fin, n):
    r = []
    for _ in range(n):
        d = debut + dt.timedelta(days=random.randint(0, (fin - debut).days))
        lib, mcc, cat, (a, b) = random.choice(COMMERCANTS)
        r.append((d, -round(random.uniform(a, b), 2), lib, mcc, cat))
    return sorted(r)

sept = achats(D("2026-09-01"), D("2026-09-28"), 34)   # facture au 30/09 : achats arrêtés deux jours ouvrés avant
octo = achats(D("2026-09-29"), D("2026-10-16"), 22)
# achat groupé au marché (deux tickets, un paiement) et achat de la facturette
octo.append((D("2026-10-10"), -72.00, "SUMUP *MARCHE DE VILLE", "5499", "Alimentation"))
octo.append((D("2026-10-12"), -86.00, "SALON LUCIE VILLE", "7230", "Divers"))
octo.sort()
cartes = {"X111": [], "X222": []}
for i, a in enumerate(octo): cartes["X222" if ("SALON LUCIE" in a[2] or "MARCHE DE VILLE" in a[2] or i % 2 == 0) else "X111"].append(a)
cartes["X222"] = sorted(cartes["X222"] + [a for a in [] ])
facture_sept = round(-sum(a[1] for a in sept), 2)

# ---------------------------------------------------------------- compte joint
lignes = []
def ligne(d, m, typ, lib1="", lib2="", mcc="", cat=""):
    lignes.append((d, m, typ, lib1, lib2, mcc, cat))
for mois in ("2026-09", "2026-10"):
    ligne(D(mois + "-01"), -3200.00, "Prélèvement", "PRLV SEPA GERANCE EXEMPLE")
    ligne(D(mois + "-03"), -182.40, "Prélèvement", "PRLV SEPA EDF clients particuli")
    ligne(D(mois + "-09"), 4100.00, "Virement", "", "VIREMENT CAISSE RETRAITE EXEMPLE")
    ligne(D(mois + "-09"), 1850.00, "Virement", "", "VIREMENT CAISSE COMPLEMENTAIRE EX")
    ligne(D(mois + "-15"), -96.99, "Prélèvement", "PRLV SEPA ORANGE SA")
ligne(D("2026-09-22"), -151.20, "Prélèvement", "PRLV SEPA EDF clients particuli")
ligne(D("2026-09-12"), -150.00, "Chèque", "CHEQUE 1234567")
ligne(D("2026-09-18"), -2000.00, "Virement", "VIR SEPA VERS LIVRET A EXEMPLE")
ligne(D("2026-09-26"), -1234.56, "Prélèvement", "PRLV SEPA American Express Cart")
ligne(D("2026-10-05"), -62.30, "Prélèvement", "PRLV SEPA EAUX DE LA VALLEE")
for a in sept + octo:  # lignes carte : doublées par la banque, sans effet sur le solde
    for _ in range(2): ligne(a[0], a[1], "Carte", "", a[2], a[3], a[4])
lignes.sort(key=lambda x: x[0])
solde0 = 85000.00
solde1 = round(solde0 + sum(l[1] for l in lignes if l[2] != "Carte") - facture_sept, 2)
with open(SORTIE / f"T_cpte_{COMPTE.replace(' ', '_')}_du_01-09-2026_au_16-10-2026.csv", "w", encoding="utf-8-sig", newline="") as f:
    f.write(f"01/09/2026;{mt(solde0)};;{COMPTE}\r\n")
    for d, m, typ, l1, l2, mcc, cat in lignes:
        f.write(f"{fr(d)};{mt(m)};{typ};;{l1};{l2};{mcc};{cat}\r\n")
    f.write(f"16/10/2026;{mt(solde1)};;{COMPTE}\r\n")

# ---------------------------------------------------------------- exports « encours » par carte
for carte, nom in (("X111", "Carte_Platinum+"), ("X222", "Carte_Premier")):
    ls = sorted(cartes[carte])
    with open(SORTIE / f"T_cpte_{COMPTE.replace(' ', '_')}_carte_{nom}_{carte}_du_29-09-2026_au_16-10-2026.csv", "w", encoding="utf-8-sig", newline="") as f:
        f.write(f"29/09/2026;{mt(round(-sum(a[1] for a in ls), 2))};;{COMPTE}\r\n")
        for d, m, lib, mcc, cat in ls: f.write(f"{fr(d)};{mt(m)};Carte;;{lib};;{mcc};{cat}\r\n")
        f.write(f"16/10/2026;0;;{COMPTE}\r\n")

# ---------------------------------------------------------------- Amex (CSV, dates mm/jj/aaaa, débits positifs)
aout = [(D("2026-08-03"), 412.30, "RESTAURANT LE GRAND EXEMPLE VILLE"), (D("2026-08-11"), 89.90, "LIBRAIRIE EN LIGNE EXEMPLE"),
        (D("2026-08-19"), 515.36, "HOTEL DE LA PLAGE EXEMPLE"), (D("2026-08-25"), 217.00, "GRAND MAGASIN EXEMPLE VILLE")]
assert abs(sum(x[1] for x in aout) - 1234.56) < 0.001
amex = aout + [(D("2026-09-26"), -1234.56, "PRELEVEMENT AUTOMATIQUE ENREGISTRE"), (D("2026-09-14"), 64.00, "TAXI EXEMPLE VILLE"),
               (D("2026-10-02"), 156.80, "RESTAURANT LE GRAND EXEMPLE VILLE"), (D("2026-10-08"), 12.99, "ABONNEMENT NUMERIQUE EXEMPLE")]
with open(SORTIE / "activity.csv", "w", encoding="cp1252", newline="") as f:
    f.write("Date,Description,Montant\r\n")
    for d, m, lib in sorted(amex, reverse=True): f.write(f'{d.strftime("%m/%d/%Y")},{lib},"{f"{m:.2f}".replace(".", ",")}"\r\n')

# ---------------------------------------------------------------- PDF : facture des cartes, relevés de placements
def pdf(nom, lignes_texte, mot_de_passe=None):
    c = canvas.Canvas(str(SORTIE / nom), pagesize=A4)
    if mot_de_passe:
        from reportlab.lib import pdfencrypt
        c = canvas.Canvas(str(SORTIE / nom), pagesize=A4, encrypt=pdfencrypt.StandardEncryption(mot_de_passe, canPrint=1))
    y = 800
    for t in lignes_texte:
        c.setFont("Helvetica", 10); c.drawString(50, y, t); y -= 16
    c.save()
pdf("CB_FACTURE_RETRAIT_DT_DIFFERE_30-09-2026.pdf", ["Détails de l'opération", "Titulaire M EXEMPLE OU Mme EXEMPLE", f"Compte {COMPTE} Compte de dépôts",
    f"Montant -{f'{facture_sept:,.2f}'.replace(',', ' ').replace('.', ',')} €", "Détails RELEVE CB AU 30/09/26", "Date de l'opération 30/09/2026", "Date de valeur 30/09/2026",
    "Nature de l'opération CB FACTURE/RETRAIT DT DIFFERE"])
pdf("LIVRETA_09999099999_20260130.pdf", ["ECRITURES DE LA PERIODE", "du 01.01.2026 au 30.01.2026 - N° 1", "DATE LIBELLE VALEUR DEBIT CREDIT", "ANCIEN SOLDE 22 950,00",
    "TOTAUX 22 950,00", "SOLDE EN EUROS 22 950,00"])
pdf("RELEVETRIMESTRIELASSURANCEVIE_20260720.pdf", ["Votre contrat", "CONTRAT EXEMPLE VIE", "Assureur : Assureur Exemple", "N° Adhésion : 999 000000001C",
    "Votre relevé trimestriel au 30 juin 2026", "Vous trouverez ci-dessous le relevé trimestriel de votre adhésion au contrat CONTRAT EXEMPLE VIE formule \"Premium\".",
    "La valeur de votre adhésion s'élève à 512 300,00 euros."])
pdf("RETRAITE_EXEMPLE_2025.pdf", ["Votre numéro de contrat", "900000001Y", "Nous avons le plaisir de vous adresser le relevé de situation 2025 de votre adhésion EXEMPLE RETRAITE. Vous trouverez",
    "SITUATION DE VOTRE CONTRAT AU 31.12.2025", "Montant de la rente viagère annuelle non réversible minimum garantie, payable", "par trimestre si vous choisissez de percevoir une retraite au terme (1)",
    "6 120,00", "Valeur de transfert (2) 141 800,00"])
pdf("PROTEGE_EXEMPLE.pdf", ["Votre numéro de contrat", "900000002F", "le relevé de situation 2025 de votre adhésion EPARGNE EXEMPLE. Vous trouverez",
    "SITUATION DE VOTRE CONTRAT AU 31.12.2025", "Epargne disponible (valeur de rachat) 73 100,00", "Cotisation MENSUELLE 200,00"], mot_de_passe="exemple")

# ---------------------------------------------------------------- tickets photographiés (images)
def ticket(nom, lignes_texte, largeur=640):
    f = ImageFont.truetype("C:/Windows/Fonts/cour.ttf", 30) if pathlib.Path("C:/Windows/Fonts/cour.ttf").exists() else ImageFont.load_default()
    h = 60 + 46 * len(lignes_texte)
    im = Image.new("RGB", (largeur, h), (250, 248, 240)); d = ImageDraw.Draw(im)
    for i, t in enumerate(lignes_texte): d.text((30, 30 + 46 * i), t, fill=(20, 20, 20), font=f)
    im.save(SORTIE / nom, quality=90)
ticket("ticket-boucherie.jpg", ["BOUCHERIE DU MARCHE", "SARL VIANDES DU MARCHE", "12 RUE EXEMPLE VILLE", "Gigot 1,2 kg      38,40", "Saucisses x6      10,20", "TOTAL TTC         48,60", "CB                48,60", "10/10/2026  10:48", "SIRET 123 456 789 00012"])
ticket("ticket-fromagerie.jpg", ["FROMAGERIE PERRET", "EURL FROMAGES PERRET", "Comte 24 mois     14,90", "Chevres x2         8,50", "TOTAL TTC         23,40", "10/10/2026  10:55"])
ticket("facturette-coiffeur.jpg", ["TICKET CLIENT", "A CONSERVER", "SALON LUCIE", "VILLE", "CARTE BANCAIRE", "XXXXXXXXXXXX0222", "SANS CONTACT", "MONTANT 86,00 EUR", "DEBIT", "LE 12/10/26 A 09:12"])

# ---------------------------------------------------------------- règles de classement fictives
import json
regles = {"format": "budget-foyer-regles", "v": 1, "lieux": ["Ville", "Campagne", "Autres"], "regles": [
    {"nom": "Loyer", "contient": "GERANCE EXEMPLE", "categorie": "Logement", "sousCategorie": "Loyer", "nature": "Courante", "lieu": "Ville", "priorite": 90, "sens": "debit"},
    {"nom": "Électricité, ville (début de mois)", "contient": "EDF", "jourMax": 12, "categorie": "Logement", "sousCategorie": "Électricité – Ville", "nature": "Courante", "lieu": "Ville", "priorite": 80, "sens": "debit"},
    {"nom": "Électricité, campagne (fin de mois)", "contient": "EDF", "jourMin": 13, "categorie": "Logement", "sousCategorie": "Électricité – Campagne", "nature": "Courante", "lieu": "Campagne", "priorite": 80, "sens": "debit"},
    {"nom": "Eau", "contient": "EAUX DE LA VALLEE", "categorie": "Logement", "sousCategorie": "Eau", "lieu": "Campagne", "priorite": 70},
    {"nom": "Téléphonie", "contient": "ORANGE", "categorie": "Télécom", "sousCategorie": "Box et mobiles", "lieu": "Autres", "priorite": 60},
    {"nom": "Pensions", "contient": "CAISSE RETRAITE|CAISSE COMPLEMENTAIRE", "categorie": "Revenus", "sousCategorie": "Pensions", "nature": "Revenu", "priorite": 90, "sens": "credit"},
    {"nom": "Paiement Amex", "contient": "AMERICAN EXPRESS|PRELEVEMENT AUTOMATIQUE", "categorie": "Transfert", "sousCategorie": "Paiement Amex", "nature": "Hors budget", "priorite": 95},
    {"nom": "Versement livret", "contient": "LIVRET A", "categorie": "Épargne", "sousCategorie": "Versements", "nature": "Hors budget", "priorite": 85}],
  "commercants": [{"nom": "Boucherie du Marché", "alias": ["BOUCHERIE DU MARCHE VILLE", "SARL VIANDES DU MARCHE"], "categorie": "Alimentation", "sousCategorie": "Courses", "lieu": "Ville"}]}
(SORTIE / "regles-fictives.json").write_text(json.dumps(regles, ensure_ascii=False, indent=1), encoding="utf-8")

# ---------------------------------------------------------------- classeur d'inventaire fictif (format v7)
wb = Workbook(); syn = wb.active; syn.title = "Synthèse"
syn.append(["Dépenses sur douze mois – base du budget 2026 (FICTIF)"]); syn.append(["Données inventées"]); syn.append([])
syn.append(["Nature", "Catégorie", "Sous-catégorie", "LCL compte", "LCL cartes (11 mois)", "LCL cartes annualisées", "Amex", "Total annuel", "Par mois", "Pertinent 2026", "Base 2026 (hors « Non »)", "Base 2026 par mois"])
budget = [("Courante", "Alimentation", "Courses", 0, 9000, 1000), ("Courante", "Alimentation", "Restaurants et cafés", 0, 3000, 2400), ("Courante", "Logement", "Loyer", 38400, 0, 0),
          ("Courante", "Logement", "Électricité – Ville", 2200, 0, 0), ("Courante", "Logement", "Électricité – Campagne", 1800, 0, 0), ("Courante", "Logement", "Eau", 750, 0, 0),
          ("Courante", "Télécom", "Box et mobiles", 1160, 0, 0), ("Courante", "Santé", "Soins et pharmacie", 0, 900, 0), ("Courante", "Culture et loisirs", "Livres, presse, sorties, jeux", 0, 600, 300),
          ("Courante", "Voiture", "Carburant et recharge", 0, 1200, 0), ("Courante", "Soins personnels", "Beauté, parfumerie, coiffeur", 0, 900, 0), ("Ponctuelle", "Voyages", "Hôtels", 0, 0, 1500)]
total = 0
for nat, cat, sous, c1, c2, c3 in budget:
    t = c1 + c2 + c3; total += t
    syn.append([nat, cat, sous, c1, c2 / 1.0909, c2, c3, t, t / 12, "Oui", t, t / 12])
syn.append(["Nouvelle", "Santé", "Mutuelle (exemple)", None, None, None, None, None, None, None, 2400, 200]); total += 2400
syn.append(["Base 2026 complète (dépenses reconduites + nouvelles)", None, None, None, None, None, None, None, None, None, total, total / 12])
ana = wb.create_sheet("Analytique")
ana.append(["Présentation analytique (FICTIF)"]); ana.append(["…"]); ana.append([])
ana.append(["Analytique", "Catégorie", "Détail", "Ville", "Campagne", "Autres", "Total", "Par mois"])
parts = {"Alimentation": (0.55, 0.40, 0.05), "Logement": (0.80, 0.18, 0.02), "Télécom": (0, 0, 1), "Santé": (0.5, 0.5, 0), "Culture et loisirs": (0.7, 0.2, 0.1), "Voiture": (0.2, 0.5, 0.3),
         "Soins personnels": (0.6, 0.4, 0), "Voyages": (0, 0, 1)}
for cat, (a, b, c) in parts.items():
    t = sum(x[3] + x[4] + x[5] for x in budget if x[1] == cat); ana.append(["Reconduite", cat, None, t * a, t * b, t * c, t, t / 12])
ana.append(["Nouvelle", "Santé", "Mutuelle (exemple)", 0, 0, 2400, 2400, 200])
ops = wb.create_sheet("Opérations")
ops.append(["Date", "Source", "Libellé", "Détail (relevé PDF)", "Montant bancaire", "Dépense nette", "Catégorie", "Sous-catégorie", "Lieu", "Ville du commerçant", "Nature", "Pertinent 2026", "Pro ?", "Référence", "Question", "Mois", "Ville / Campagne / Autres"])
classement = {"BOUCHERIE DU MARCHE VILLE": ("Alimentation", "Courses", "Ville"), "PRIMEUR DE LA PLACE VILLE": ("Alimentation", "Courses", "Ville"), "SUPERMARCHE EXEMPLE CAMPAGNE": ("Alimentation", "Courses", "Campagne"),
              "PHARMACIE CENTRALE VILLE": ("Santé", "Soins et pharmacie", "Ville"), "LIBRAIRIE DES ECOLES VILLE": ("Culture et loisirs", "Livres, presse, sorties, jeux", "Ville"),
              "BRASSERIE DU COIN VILLE": ("Alimentation", "Restaurants et cafés", "Ville"), "STATION SERVICE CAMPAGNE": ("Voiture", "Carburant et recharge", "Campagne"), "FROMAGERIE PERRET VILLE": ("Alimentation", "Courses", "Ville")}
for d, m, lib, mcc, cat in achats(D("2026-01-02"), D("2026-08-27"), 180):
    c, s, l = classement[lib]
    ops.append([dt.datetime(d.year, d.month, d.day), "LCL carte", lib, None, m, -m, c, s, None, None, "Courante", "Oui", None, None, None, dt.datetime(d.year, d.month, 1), l])
for mois in range(1, 9):
    d = dt.datetime(2026, mois, 1)
    ops.append([d, "LCL compte", "PRLV SEPA GERANCE EXEMPLE", None, -3200, 3200, "Logement", "Loyer", "Ville", None, "Courante", "Oui", None, None, None, d, "Ville"])
    d9 = dt.datetime(2026, mois, 9)
    ops.append([d9, "LCL compte", "VIREMENT CAISSE RETRAITE EXEMPLE", None, 4100, -4100, "Revenus", "Pensions", None, None, "Revenu", "—", None, None, None, d, "Autres"])
par = wb.create_sheet("Paramètres"); par.append(["Paramètres"]); par.append(["• Compte de dépôts " + COMPTE + " : export CSV (FICTIF)"])
pro = wb.create_sheet("Projection épargne")
for r in [["Projection (FICTIF)"], ["Hypothèses"], ["Pensions brutes annuelles du foyer (2026)", 72000], ["Prélèvements sociaux sur les pensions", 0.1],
          ["Voiture : solde à la livraison (2026)", 30000], ["Travaux de la maison (2026)", 40000], ["Loyer actuel, charges comprises (par an)", 38400]]: pro.append(r)
wb.save(SORTIE / "Dépenses fictives - inventaire v7.xlsx")
print("Jeu fictif écrit dans", SORTIE, "· facture des cartes de septembre :", facture_sept, "· solde final :", solde1)
